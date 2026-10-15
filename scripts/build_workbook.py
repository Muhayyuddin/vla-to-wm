"""Build data/companion_workbook.xlsx from the CSV/JSON files in data/.

Sheets: Corpus (the 144 included reports), Evidence ledger (Table 3),
Benchmarks (Table 6), LIBERO (Table 7), Datasets (Table 9 with source
versions), Data scale (Fig. 7), World-Gymnast (Fig. 12), Screening (Fig. 3).
"""
from pathlib import Path
import csv, json
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / 'data'

BENCHMARKS = [  # Table 6 of the paper
    ('VLA and policy benchmarks', 'LIBERO', 2023, 'S', 'S/L', '130 language-conditioned tasks in four suites', 'High reported scores; sensitivity to perturbations'),
    ('VLA and policy benchmarks', 'LIBERO-Plus', 2025, 'S', 'S/L', 'Robustness across seven perturbation axes', 'Restricted to the scope of the underlying LIBERO tasks'),
    ('VLA and policy benchmarks', 'SimplerEnv', 2024, 'S/R', 'S', 'Control-matched simulation of two physical platforms', 'Limited embodiment coverage'),
    ('VLA and policy benchmarks', 'CALVIN', 2022, 'S', 'L', 'Long-horizon sequences and ABC->D generalization', 'Single simulated robot platform'),
    ('VLA and policy benchmarks', 'RLBench', 2020, 'S', 'S/L', 'One hundred manipulation tasks', 'Dependence on scripted demonstrations'),
    ('VLA and policy benchmarks', 'ManiSkill3', 2024, 'S', 'S/L', 'GPU-parallel large-scale evaluation', 'Simulation-only physics'),
    ('VLA and policy benchmarks', 'RoboCasa', 2024, 'S', 'L', 'Generated kitchen environments', 'Simulation and layout bias'),
    ('VLA and policy benchmarks', 'Meta-World', 2019, 'S', 'S', 'Multi-task and meta-reinforcement learning', 'Limited visual diversity and saturation on some tasks'),
    ('VLA and policy benchmarks', 'RoboArena', 2025, 'R', 'S/L', 'Distributed real-robot A/B evaluation', 'Throughput and variation across sites'),
    ('VLA and policy benchmarks', 'OXE, DROID, BridgeData V2, AgiBot World', '2023-25', 'R', '-', 'Training datasets with held-out evaluation splits', 'Not controlled benchmarks'),
    ('World-model benchmarks', 'WorldModelBench', 2025, 'S/R', 'S', 'Instruction-conditioned generation and physical consistency', 'Evaluates generation rather than control'),
    ('World-model benchmarks', 'Physics-IQ', 2025, 'R', 'S', 'Physical understanding from paired videos', 'No action conditioning'),
    ('World-model benchmarks', 'VBench-2.0', 2025, 'S/R', 'S', 'Intrinsic faithfulness, including physical behavior', 'Targets general video rather than robotics'),
    ('World-model benchmarks', 'EWMBench', 2025, 'S/R', 'S/L', 'Scene, motion, and semantic quality in embodied prediction', 'Closely associated with one system family'),
    ('World-model benchmarks', 'WorldArena', 2026, 'S/R', 'S/L', 'Perceptual and functional evaluation', 'Recent benchmark with a limited task source'),
    ('World-model benchmarks', 'WorldArena 2.0', 2026, 'S/R', 'S/L', 'Visuotactile prediction, interactive learning, and multi-embodiment evaluation', 'Recent preprint without broad independent replication'),
    ('World-model benchmarks', 'ManipArena', 2026, 'S/R', 'S/L', 'Shared evaluation of VLAs and world models using real-to-simulation twins', 'Recent benchmark restricted to 20 tasks'),
]
LIBERO = [  # Table 7 of the paper (success rate, %)
    ('OpenVLA', 'Published baseline', 84.7, 88.4, 79.2, 53.7, 76.5),
    ('OpenVLA-OFT', 'Filtered training set; third-person image and language', 96.2, 98.3, 96.2, 90.7, 95.3),
    ('OpenVLA-OFT', 'Filtered training set; third-person and wrist images, optional proprioception, and language', 97.6, 98.4, 97.9, 94.5, 97.1),
    ('pi_0', 'Result reported by the OFT study using its modified training and observation configuration', 96.8, 98.8, 95.8, 85.2, 94.2),
    ('SimpleVLA-RL', 'Full-trajectory supervised and reinforcement-learning configuration; all four suites', 99.4, 99.1, 99.2, 98.5, 99.1),
    ('Cosmos Policy', 'Direct parallel decoding; filtered policy data and unfiltered world/value training data; 50 episodes per task over three seeds', 98.1, 100.0, 98.2, 97.6, 98.5),
]
DATASETS = [  # Table 9 of the paper, with source version and verification date where recorded
    ('Open X-Embodiment', '970k+ trajectories, 22 embodiments', 'Robot trajectories', 'Per-dataset (mixed)', 'ICRA 2024, doi:10.1109/ICRA57147.2024.10611477', 'https://robotics-transformer-x.github.io/', '2026-10-15'),
    ('DROID', '76k episodes, 564 scenes', 'Robot trajectories', 'CC-BY-4.0', 'arXiv:2403.12945v2 (RSS 2024)', 'https://arxiv.org/abs/2403.12945', '2026-08-26'),
    ('BridgeData V2', '60k+ trajectories', 'Robot trajectories', 'CC-BY', 'CoRL 2023, PMLR v229', 'https://proceedings.mlr.press/v229/walke23a.html', '2026-10-15'),
    ('AgiBot World', '1M+ trajectories, 100+ robots', 'Robot trajectories', 'CC-BY-NC-SA-4.0', 'IROS 2025; arXiv:2503.06669', 'https://arxiv.org/abs/2503.06669', '2026-10-15'),
    ('LeRobot', 'Community aggregation (growing)', 'Robot trajectories', 'Apache-2.0 (code); per-dataset licenses vary', 'arXiv:2602.22818', 'https://arxiv.org/abs/2602.22818', '2026-10-15'),
    ('Ego4D', '3,670 h egocentric video', 'Human video', 'Ego4D license', 'CVPR 2022', 'https://ego4d-data.org/', '2026-10-15'),
    ('VideoMix22M', '22M videos / 1M+ h', 'Internet video', 'Per-source (mixed)', 'arXiv:2506.09985v1 (V-JEPA 2)', 'https://arxiv.org/abs/2506.09985', '2026-08-26'),
    ('DreamGen trajectories', 'Synthetic neural trajectories', 'Generated', 'Apache-2.0', 'CoRL 2025, PMLR v305', 'https://proceedings.mlr.press/v305/jang25a.html', '2026-10-15'),
]

def sheet(wb, title, header, rows):
    ws = wb.create_sheet(title); ws.append(header)
    for c in ws[1]: c.font = Font(bold=True)
    for r in rows: ws.append(list(r))
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(12, len(str(h)) + 2))
    ws.freeze_panes = 'A2'

def main():
    wb = Workbook(); wb.remove(wb.active)
    reg = list(csv.DictReader(open(DATA / 'literature_register.csv', encoding='utf-8')))
    inc = [r for r in reg if r['decision'] == 'Include']
    sheet(wb, 'Corpus', list(inc[0].keys()), [r.values() for r in inc])
    led = list(csv.DictReader(open(DATA / 'evidence_ledger.csv', encoding='utf-8')))
    sheet(wb, 'Evidence ledger', list(led[0].keys()), [r.values() for r in led])
    sheet(wb, 'Benchmarks', ['Group', 'Benchmark', 'Year', 'Type (S/R)', 'Horizon', 'Primary focus', 'Principal caveat'], BENCHMARKS)
    sheet(wb, 'LIBERO', ['Model', 'Reported configuration', 'Spatial', 'Object', 'Goal', 'Long', 'Avg'], LIBERO)
    sheet(wb, 'Datasets', ['Corpus', 'Reported scale', 'Modality', 'License', 'Source version', 'Source URL', 'Verification date'], DATASETS)
    ds = list(csv.DictReader(open(DATA / 'figure07_data_scale.csv', encoding='utf-8')))
    sheet(wb, 'Data scale (Fig. 7)', list(ds[0].keys()), [r.values() for r in ds])
    wg = list(csv.DictReader(open(DATA / 'figure12_world_gymnast.csv', encoding='utf-8')))
    sheet(wb, 'World-Gymnast (Fig. 12)', list(wg[0].keys()), [r.values() for r in wg])
    sc = json.loads((DATA / 'screening_counts.json').read_text())
    flat = []
    def walk(d, prefix=''):
        for k, v in d.items():
            if isinstance(v, dict): walk(v, prefix + k + ' / ')
            else: flat.append((prefix + k, v))
    walk(sc); sheet(wb, 'Screening (Fig. 3)', ['Quantity', 'Value'], flat)
    wb.save(DATA / 'companion_workbook.xlsx'); print('wrote', DATA / 'companion_workbook.xlsx')

if __name__ == '__main__':
    main()
